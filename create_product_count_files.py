from pathlib import Path
import csv
from collections import defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo

out = Path('/home/ubuntu')
# A compact Superstore-style dataset with repeated Product IDs to demonstrate unique counting.
products = [
    ('FUR-001','Office Chair','Furniture'), ('FUR-002','Writing Desk','Furniture'),
    ('FUR-003','Bookcase','Furniture'), ('FUR-004','Conference Table','Furniture'),
    ('FUR-005','Filing Cabinet','Furniture'), ('FUR-006','Storage Cabinet','Furniture'),
    ('FUR-007','Standing Desk','Furniture'),
    ('OFF-001','Paper','Office Supplies'), ('OFF-002','Stapler','Office Supplies'),
    ('OFF-003','Pens','Office Supplies'), ('OFF-004','Binders','Office Supplies'),
    ('OFF-005','Envelopes','Office Supplies'), ('OFF-006','Labels','Office Supplies'),
    ('OFF-007','Scissors','Office Supplies'), ('OFF-008','Storage Box','Office Supplies'),
    ('OFF-009','Paper Clips','Office Supplies'), ('OFF-010','Notepad','Office Supplies'),
    ('TEC-001','Laptop','Technology'), ('TEC-002','Monitor','Technology'),
    ('TEC-003','Keyboard','Technology'), ('TEC-004','Mouse','Technology'),
    ('TEC-005','Printer','Technology'), ('TEC-006','Tablet','Technology'),
]
rows=[]
for i in range(60):
    pid, name, cat = products[(i*3) % len(products)]
    rows.append({'Order ID':f'ORD-{i+1:03d}', 'Product ID':pid, 'Product Name':name,
                 'Category':cat, 'Quantity':(i%4)+1, 'Sales':round(100+(i%10)*37.5,2)})

unique=defaultdict(set)
for r in rows: unique[r['Category']].add(r['Product ID'])
counts={c:len(unique[c]) for c in ['Furniture','Office Supplies','Technology']}
top=max(counts, key=counts.get); highest=counts[top]

raw_csv=out/'Superstore_Product_Dataset.csv'
with raw_csv.open('w',newline='',encoding='utf-8') as f:
    w=csv.DictWriter(f,fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
summary_csv=out/'Product_Count_Summary.csv'
with summary_csv.open('w',newline='',encoding='utf-8') as f:
    w=csv.writer(f); w.writerow(['Category','Unique Product Count'])
    for c in counts: w.writerow([c,counts[c]])
    w.writerow([]); w.writerow(['Top Category',top]); w.writerow(['Highest Product Count',highest])

wb=Workbook(); raw=wb.active; raw.title='Raw Data'; s=wb.create_sheet('Product Count Analysis')
headers=list(rows[0]); raw.append(headers)
for r in rows: raw.append([r[h] for h in headers])
blue='1F4E78'; light='D9EAF7'; orange='F4B183'; green='70AD47'; thin=Side(style='thin',color='B7C9D6')
for cell in raw[1]: cell.fill=PatternFill('solid',fgColor=blue); cell.font=Font(color='FFFFFF',bold=True); cell.alignment=Alignment(horizontal='center')
for row in raw.iter_rows(min_row=2):
    for cell in row: cell.border=Border(bottom=thin)
for col,width in {'A':14,'B':15,'C':22,'D':20,'E':12,'F':14}.items(): raw.column_dimensions[col].width=width
raw.freeze_panes='A2'; raw.auto_filter.ref=raw.dimensions
style=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True,showColumnStripes=False)
tab=Table(displayName='SuperstoreData',ref=f'A1:F{raw.max_row}'); tab.tableStyleInfo=style; raw.add_table(tab)

s.merge_cells('A1:D1'); s['A1']='PRODUCT COUNT ANALYSIS'; s['A1'].font=Font(size=18,bold=True,color='FFFFFF'); s['A1'].fill=PatternFill('solid',fgColor='17365D'); s['A1'].alignment=Alignment(horizontal='center')
for c in range(1,5): s.cell(1,c).fill=PatternFill('solid',fgColor='17365D')
s['A3']='Category'; s['B3']='Unique Product Count'; s['C3']='Formula / Check'
for x in ['A3','B3','C3']:
    s[x].font=Font(bold=True,color='FFFFFF'); s[x].fill=PatternFill('solid',fgColor=blue); s[x].alignment=Alignment(horizontal='center')
for i,c in enumerate(['Furniture','Office Supplies','Technology'],4):
    s[f'A{i}']=c
    # Modern Excel formula for unique products by category.
    s[f'B{i}']=f'=COUNTA(UNIQUE(FILTER(\'Raw Data\'!$B$2:$B$61,\'Raw Data\'!$D$2:$D$61=A{i})))'
    s[f'C{i}']=f'Expected result: {counts[c]}'
    s[f'B{i}'].number_format='0'
for row in s.iter_rows(min_row=3,max_row=6,min_col=1,max_col=3):
    for cell in row: cell.border=Border(left=thin,right=thin,top=thin,bottom=thin)
s['A9']='Top Category'; s['B9']='=INDEX(A4:A6,MATCH(MAX(B4:B6),B4:B6,0))'; s['C9']=top
s['A10']='Highest Product Count'; s['B10']='=MAX(B4:B6)'; s['C10']=highest
for r in [9,10]:
    for c in range(1,4): s.cell(r,c).border=Border(left=thin,right=thin,top=thin,bottom=thin)
    s[f'A{r}'].font=Font(bold=True); s[f'A{r}'].fill=PatternFill('solid',fgColor=light)
s['A13']='Simple Explanation'; s['A13'].font=Font(bold=True,size=13,color='17365D')
s['A14']='The table counts each Product ID only once within each category.'
s['A15']=f'Top Category: {top}'
s['A16']=f'Highest Product Count: {highest}'
for col,width in {'A':24,'B':24,'C':28,'D':4}.items(): s.column_dimensions[col].width=width
s.freeze_panes='A4'; s.sheet_view.showGridLines=False; raw.sheet_view.showGridLines=False
chart=BarChart(); chart.type='col'; chart.style=10; chart.title='Unique Product Count by Category'; chart.y_axis.title='Unique Products'; chart.x_axis.title='Category'
data=Reference(s,min_col=2,min_row=3,max_row=6); cats=Reference(s,min_col=1,min_row=4,max_row=6); chart.add_data(data,titles_from_data=True); chart.set_categories(cats); chart.height=7; chart.width=12; s.add_chart(chart,'E3')

xlsx=out/'Superstore_Product_Count_Analysis.xlsx'; wb.save(xlsx)
# Verify files and formulas
check=load_workbook(xlsx,data_only=False); assert check.sheetnames==['Raw Data','Product Count Analysis']; assert check['Product Count Analysis']['B9'].value=='=INDEX(A4:A6,MATCH(MAX(B4:B6),B4:B6,0))'
print(f'Created: {xlsx}'); print(f'Created: {raw_csv}'); print(f'Created: {summary_csv}'); print(f'Unique counts: {counts}; Top category: {top} ({highest})')
